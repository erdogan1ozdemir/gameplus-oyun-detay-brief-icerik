#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Oyun detay sayfası içeriğini Word dosyasına basar.

Kullanım:
    python3 icerik_docx.py --json icerik.json --out battlefield-6-icerik.docx

icerik.json biçimi:
{
  "oyun": "Battlefield 6",
  "meta": "Gövde 827 kelime · 5 H2, 4 H3 · 2 tablo · 6 iç link · 8 SSS",
  "linkler": {"LINK1": ["anchor metni", "https://gameplus.com.tr/..."]},
  "govde": [["H2","Başlık"], ["p","Paragraf, içinde [LINK1] geçebilir"],
            ["H3","Alt başlık"], ["li","Numaralı madde"], ["mad","Madde imli madde"],
            ["tablo", [["Sütun","Sütun"],["satır","satır"]]]],
  "sss": [["Soru?","Yanıt"]],
  "not": "Kaynak ve karar notları"
}
Gövde H2 ile başlar; sayfada H1 oyun adı olarak bulunduğu için belgeye H1 yazılmaz.
"""
import argparse, json, re

TEAL, GRI, CORAL, LINK, FN = "10332F", "E0E0E0", "FF7B52", "0B5FB0", "Calibri"


def renk(run, hex_kod):
    from docx.shared import RGBColor
    run.font.color.rgb = RGBColor.from_string(hex_kod)


def koprü(p, metin, url):
    """python-docx'te yerleşik köprü yok; ilişki kaydı elle eklenir."""
    from docx.oxml.shared import OxmlElement, qn
    r_id = p.part.relate_to(url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True)
    h = OxmlElement("w:hyperlink"); h.set(qn("r:id"), r_id)
    r = OxmlElement("w:r"); rPr = OxmlElement("w:rPr")
    for etiket, deger in (("w:color", LINK), ("w:u", "single")):
        e = OxmlElement(etiket); e.set(qn("w:val"), deger); rPr.append(e)
    rf = OxmlElement("w:rFonts"); rf.set(qn("w:ascii"), FN); rf.set(qn("w:hAnsi"), FN); rPr.append(rf)
    sz = OxmlElement("w:sz"); sz.set(qn("w:val"), "21"); rPr.append(sz)
    r.append(rPr)
    t = OxmlElement("w:t"); t.text = metin; t.set(qn("xml:space"), "preserve"); r.append(t)
    h.append(r); p._p.append(h)


def metni_bas(p, metin, linkler, boyut=10.5):
    from docx.shared import Pt
    parcalar = re.split(r"(\[LINK\d+\])", metin)
    for parca in parcalar:
        m = re.fullmatch(r"\[(LINK\d+)\]", parca)
        if m and m.group(1) in linkler:
            ad, url = linkler[m.group(1)]
            koprü(p, ad, url)
        elif parca:
            r = p.add_run(parca); r.font.name = FN; r.font.size = Pt(boyut); renk(r, TEAL)


def main():
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT

    ap = argparse.ArgumentParser()
    ap.add_argument("--json", required=True)
    ap.add_argument("--out", required=True)
    a = ap.parse_args()
    d = json.load(open(a.json, encoding="utf-8"))
    linkler = {k: tuple(v) for k, v in (d.get("linkler") or {}).items()}

    doc = Document()
    st = doc.styles["Normal"]; st.font.name = FN; st.font.size = Pt(10.5)
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Cm(2); s.left_margin = s.right_margin = Cm(2)

    p = doc.add_paragraph(); r = p.add_run(f"{d.get('oyun','Oyun')} | Oyun Detay Sayfası İçeriği")
    r.bold = True; r.font.size = Pt(15); r.font.name = FN; renk(r, TEAL)
    if d.get("meta"):
        p = doc.add_paragraph(); r = p.add_run(d["meta"])
        r.font.size = Pt(8.5); r.font.name = FN; renk(r, "5A6B68")

    for tip, icerik in d["govde"]:
        if tip in ("H2", "H3"):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14 if tip == "H2" else 11)
            p.paragraph_format.space_after = Pt(5)
            r = p.add_run(icerik); r.bold = True; r.font.name = FN
            r.font.size = Pt(13 if tip == "H2" else 11.5); renk(r, TEAL)
            p.style = doc.styles["Heading 1" if tip == "H2" else "Heading 2"]
            for run in p.runs:
                run.font.name = FN; run.bold = True
                run.font.size = Pt(13 if tip == "H2" else 11.5); renk(run, TEAL)
        elif tip == "tablo":
            t = doc.add_table(rows=0, cols=len(icerik[0])); t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.LEFT
            for i, satir in enumerate(icerik):
                hucreler = t.add_row().cells
                for j, deger in enumerate(satir):
                    hucreler[j].text = ""
                    par = hucreler[j].paragraphs[0]
                    par.alignment = WD_ALIGN_PARAGRAPH.LEFT if (j == 0 and icerik[0][0] == "") else WD_ALIGN_PARAGRAPH.CENTER
                    rr = par.add_run(str(deger)); rr.font.name = FN; rr.font.size = Pt(9.5)
                    if i == 0:
                        rr.bold = True; renk(rr, "FFFFFF")
                        from docx.oxml.shared import OxmlElement, qn
                        sh = OxmlElement("w:shd"); sh.set(qn("w:fill"), TEAL)
                        hucreler[j]._tc.get_or_add_tcPr().append(sh)
                    else:
                        renk(rr, TEAL)
            doc.add_paragraph()
        elif tip in ("li", "mad"):
            p = doc.add_paragraph(style="List Number" if tip == "li" else "List Bullet")
            p.paragraph_format.space_after = Pt(3)
            metni_bas(p, icerik, linkler)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(7)
            metni_bas(p, icerik, linkler)

    if d.get("sss"):
        p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(20)
        r = p.add_run("Sıkça Sorulan Sorular"); r.bold = True; r.font.size = Pt(13); r.font.name = FN; renk(r, TEAL)
        for soru, yanit in d["sss"]:
            p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(9); p.paragraph_format.space_after = Pt(2)
            r = p.add_run(soru); r.bold = True; r.font.size = Pt(11); r.font.name = FN; renk(r, TEAL)
            p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(4)
            r = p.add_run(yanit); r.font.size = Pt(10.5); r.font.name = FN; renk(r, TEAL)

    if d.get("not"):
        p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(18)
        r = p.add_run("Not: "); r.bold = True; r.font.size = Pt(8.5); r.font.name = FN; renk(r, CORAL)
        r = p.add_run(d["not"]); r.font.size = Pt(8.5); r.font.name = FN; renk(r, "5A6B68")

    doc.save(a.out)
    kelime = sum(len(re.sub(r"\[LINK\d+\]", "x", i).split())
                 for t, i in d["govde"] if t in ("p", "li", "mad"))
    print(f"yazıldı: {a.out} · gövde {kelime} kelime · "
          f"{sum(1 for t,_ in d['govde'] if t=='H2')} H2 · {sum(1 for t,_ in d['govde'] if t=='H3')} H3 · "
          f"{len(d.get('sss',[]))} SSS")


if __name__ == "__main__":
    main()
