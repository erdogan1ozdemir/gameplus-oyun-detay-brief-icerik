#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""İçerik JSON'undan CMS giriş dosyası üretir: URL, Meta title, Meta description, Uzun açıklama (HTML) ve
FAQ (HTML + betik). Çıktı Word dosyasıdır; HTML kod olarak, satır satır ve eşit aralıklı yazıyla basılır.

Kullanım:
    python3 cms_docx.py --json kaynak/battlefield-6.json --out "HTML battlefield-6-oyun-detay-sayfasi-icerik.docx"
        [--onizleme onizleme-sss-battlefield-6.html] [--ad "Battlefield 6"] [--alan https://prp1.gameplus.com.tr]

Kurallar (gameplus-oyun-detay-sayfasi-kurallari.docx ve CMS örnek dosyası, 28.09.2026):
- URL: {alan}/gfn/oyunlar/oyun/{slug}. Yol CMS örneğindeki yapıdır; slug kural dokümanının slug kuralıyla
  GFN kaydının title alanından üretilir (® ™ temizlenir, & -> and, seri numarasındaki roma rakamı -> rakam,
  Türkçe harf ASCII'ye katlanır, küçük harf, tire).
- Meta title: "{Oyun adı} - GeForce NOW | Gameplus".
- Meta description: "{Oyun adı}, GeForce NOW powered by GAME+ kütüphanesinde. {Tür} türündeki oyunu bulut
  üzerinden oyna." Tür, GFN kaydındaki genres[0] alanının karşılığıdır.
- Uzun açıklama: gövdenin sade HTML'i (h2, h3, p, ul, ol, table, a, strong); belge başlığı ve SSS girmez,
  satır içi stil ve font yoktur.
- FAQ: CMS örneğindeki buton + toggleFaq yapısı korunur, görünüm blog SSS'siyle aynıdır (kenarlıklı kart,
  kalın soru, sarı nabız atan +, açılınca 45° döner). gp-content sınıfı ve <style> bloğu kullanılmaz;
  animasyon betikten çalışır. Lisanslı font gömülmez.
"""
import argparse, html, json, os, re, sys, unicodedata

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

TUR = {"ACTION": "Aksiyon", "ADVENTURE": "Macera", "INDIE": "Bağımsız", "SIMULATION": "Simülasyon",
       "STRATEGY": "Strateji", "ROLE_PLAYING": "Rol Yapma", "CASUAL": "Basit Eğlence",
       "FIRST_PERSON_SHOOTER": "FPS", "FREE_TO_PLAY": "Oynaması Ücretsiz",
       "MASSIVELY_MULTIPLAYER_ONLINE": "MMO", "RACING": "Yarış", "SPORTS": "Spor", "FAMILY": "Aile Dostu",
       "PUZZLE": "Bulmaca", "PLATFORMER": "Platform", "FIGHTING": "Dövüş Oyunu", "ARCADE": "Arcade",
       "MULTIPLAYER_ONLINE_BATTLE_ARENA": "MOBA", "DEMO": "Demo", "TECH_DEMO": "Teknolojik Demo"}
ROMA = {"II": "2", "III": "3", "IV": "4", "V": "5", "VI": "6", "VII": "7", "VIII": "8", "IX": "9", "X": "10",
        "XI": "11", "XII": "12", "XIII": "13", "XIV": "14", "XV": "15", "XVI": "16"}


def temiz_ad(title):
    return re.sub(r"\s+", " ", re.sub(r"[®™©℠​]", "", title)).strip()


def slug(ad):
    s = temiz_ad(ad)
    s = re.sub(r"\b(?:[A-Za-z]\.){2,}", lambda m: m.group(0).replace(".", ""), s)   # S.T.A.L.K.E.R. -> STALKER
    s = s.replace(".", " ").replace("&", " and ")
    s = re.sub(r"(\d),(\d{3})", r"\1\2", s)
    kel = s.split()
    for i in range(1, len(kel)):                    # ilk kelimeye ve tek başına X / XX'e uygulanmaz
        k = kel[i].strip(":,")
        if k in ROMA and k not in ("X", "XX"):
            kel[i] = kel[i].replace(k, ROMA[k])
    s = " ".join(kel).replace("²", "2").replace("³", "3")
    s = s.translate(str.maketrans("çğıİöşüÇĞÖŞÜ", "cgiIosuCGOSU"))
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode()
    return re.sub(r"-+", "-", re.sub(r"[^a-z0-9]+", "-", s.lower())).strip("-")[:64]


def satir_ici(metin, linkler):
    """**kalın** -> <strong>, [LINKn] -> <a href>. Metin önce HTML'den kaçırılır."""
    m = html.escape(metin.strip(), quote=False)
    m = re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", m)
    def link(x):
        anchor, url = linkler[x.group(1)]
        return f'<a href="{html.escape(url)}">{html.escape(anchor, quote=False)}</a>'
    m = re.sub(r"\[(LINK\d+)\]", link, m)
    assert "[LINK" not in m and "**" not in m, m[:80]
    return m


def uzun_aciklama(d):
    L, liste = ["<div>"], None
    def kapat():
        nonlocal liste
        if liste: L.append(f"  </{liste}>"); liste = None
    for tip, x in d["govde"]:
        if tip in ("mad", "li"):
            etiket = "ul" if tip == "mad" else "ol"
            if liste != etiket:
                kapat(); L.append(f"  <{etiket}>"); liste = etiket
            L.append(f"    <li>{satir_ici(x, d['linkler'])}</li>"); continue
        kapat()
        if tip in ("H2", "H3"):
            if len(L) > 1: L.append("")
            L.append(f"  <{tip.lower()}>{satir_ici(x, d['linkler'])}</{tip.lower()}>")
        elif tip == "p":
            L.append(f"  <p>{satir_ici(x, d['linkler'])}</p>")
        elif tip == "tablo":
            L.append("  <table>")
            L.append("    <thead><tr>" + "".join(f"<th>{satir_ici(c, d['linkler'])}</th>" for c in x[0]) + "</tr></thead>")
            L.append("    <tbody>")
            for r in x[1:]:
                L.append("      <tr>" + "".join(f"<td>{satir_ici(c, d['linkler'])}</td>" for c in r) + "</tr>")
            L.append("    </tbody>")
            L.append("  </table>")
        else:
            raise SystemExit(f"bilinmeyen öğe: {tip}")
    kapat(); L.append("</div>")
    return "\n".join(L)


SARI, SARI_KOYU, CIZGI = "#FFC900", "#f59e0b", "#29292b"


def faq(d):
    L = ['<div style="width:100%; background:#0d0d0d; padding:clamp(20px, 4vw, 40px); font-family:inherit; box-sizing:border-box;">',
         "",
         '  <h2 style="font-size:clamp(24px, 3.2vw, 32px); line-height:1.25; font-weight:600; color:#ffffff; margin:0 0 24px 0;">Sık Sorulan Sorular</h2>',
         "",
         '  <div id="faq-list">']
    for q, a in d["sss"]:
        L += ["",
              f'    <div class="faq-item" style="margin:0 0 10px 0; border:1px solid {CIZGI}; border-radius:10px; overflow:hidden; background:transparent; box-shadow:0 2px 8px rgba(0,0,0,0.4);">',
              '      <button type="button" onclick="toggleFaq(this)" aria-expanded="false" style="width:100%; background:none; border:none; cursor:pointer; display:flex; align-items:center; gap:10px; padding:14px 16px; text-align:left; font-family:inherit;">',
              f'        <span class="faq-icon" aria-hidden="true" style="display:inline-flex; align-items:center; justify-content:center; width:22px; height:22px; flex-shrink:0; color:{SARI}; transition:transform 0.25s ease, color 0.2s;"><span class="faq-pulse" style="display:inline-block; font-size:22px; font-weight:300; line-height:1;">+</span></span>',
              f'        <span class="faq-q" style="flex:1; font-size:clamp(16px, 1.6vw, 19px); font-weight:700; line-height:1.35; letter-spacing:-0.005em; color:#f3f4f6;">{html.escape(q.strip(), quote=False)}</span>',
              "      </button>",
              f'      <div class="faq-answer" style="display:none; padding:14px 20px 18px clamp(16px, 5vw, 48px); border-top:1px solid {CIZGI};">',
              f'        <p style="font-size:clamp(15px, 1.5vw, 17.5px); color:#B2B2B2; margin:0; line-height:1.55;">{html.escape(a.strip(), quote=False)}</p>',
              "      </div>",
              "    </div>"]
    L += ["", "  </div>", "</div>", "", "<script>",
          "function toggleFaq(btn) {",
          "  var answer = btn.nextElementSibling;",
          "  var isOpen = answer.style.display === 'block';",
          "",
          "  // Tüm açık olanları kapat",
          "  document.querySelectorAll('#faq-list .faq-answer').forEach(function(el) {",
          "    el.style.display = 'none';",
          "  });",
          "  document.querySelectorAll('#faq-list .faq-item button').forEach(function(el) {",
          "    el.setAttribute('aria-expanded', 'false');",
          "    el.querySelector('.faq-icon').style.transform = 'rotate(0deg)';",
          "  });",
          "",
          "  // Tıklananı aç (zaten açıksa kapalı kalır); + işareti 45 derece dönüp x olur",
          "  if (!isOpen) {",
          "    answer.style.display = 'block';",
          "    btn.setAttribute('aria-expanded', 'true');",
          "    btn.querySelector('.faq-icon').style.transform = 'rotate(45deg)';",
          "  }",
          "}",
          "",
          "(function () {",
          "  var sakin = window.matchMedia && window.matchMedia('(prefers-reduced-motion: reduce)').matches;",
          "  document.querySelectorAll('#faq-list .faq-item button').forEach(function(btn) {",
          "    var icon = btn.querySelector('.faq-icon');",
          f"    btn.addEventListener('mouseenter', function () {{ icon.style.color = '{SARI_KOYU}'; }});",
          f"    btn.addEventListener('mouseleave', function () {{ icon.style.color = '{SARI}'; }});",
          "    var pulse = btn.querySelector('.faq-pulse');",
          "    if (!sakin && pulse.animate) {",
          "      pulse.animate([",
          "        { transform: 'scale(1)', opacity: 1 },",
          "        { transform: 'scale(1.18)', opacity: 0.85 },",
          "        { transform: 'scale(1)', opacity: 1 }",
          "      ], { duration: 2200, iterations: Infinity, easing: 'ease-in-out' });",
          "    }",
          "  });",
          "})();",
          "</script>"]
    return "\n".join(L)


def docx_yaz(yol, alanlar):
    from docx import Document
    from docx.shared import Pt, RGBColor
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)
    for etiket, deger, kod in alanlar:
        p = doc.add_paragraph(); r = p.add_run(etiket); r.bold = True
        r.font.color.rgb = RGBColor.from_string("10332F")
        p.paragraph_format.space_before = Pt(14); p.paragraph_format.space_after = Pt(4)
        for satir in deger.split("\n"):
            q = doc.add_paragraph(); q.paragraph_format.space_after = Pt(0); q.paragraph_format.space_before = Pt(0)
            r = q.add_run(satir)
            if kod: r.font.name = "Consolas"; r.font.size = Pt(9)
    doc.save(yol)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--json", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--onizleme")
    ap.add_argument("--ad", help="GFN kaydını bulmak için oyun adı (varsayılan: JSON'daki oyun)")
    ap.add_argument("--alan", default="https://prp1.gameplus.com.tr")
    a = ap.parse_args()
    d = json.load(open(a.json, encoding="utf-8"))
    import dil_hucresi as dh
    g = dh.bul(a.ad or d["oyun"])
    if not g: raise SystemExit(f"GFN kaydı bulunamadı: {a.ad or d['oyun']}")
    ad = temiz_ad(g["title"])
    url = f"{a.alan.rstrip('/')}/gfn/oyunlar/oyun/{slug(g['title'])}"
    title = f"{ad} - GeForce NOW | Gameplus"
    tur = TUR.get((g.get("genres") or [None])[0])
    desc = f"{ad}, GeForce NOW powered by GAME+ kütüphanesinde." + (f" {tur} türündeki oyunu bulut üzerinden oyna." if tur else "")
    uzun, sss = uzun_aciklama(d), faq(d)
    for yasak in ("@font-face", "data:font", ".woff", ".otf", "gp-content", "<style"):
        assert yasak not in uzun + sss, yasak
    docx_yaz(a.out, [("URL:", url, False), ("Meta title:", title, False), ("Meta description:", desc, False),
                     ("Uzun açıklama:", uzun, True), ("FAQ:", sss, True)])
    if a.onizleme:
        open(a.onizleme, "w", encoding="utf-8").write(
            '<!doctype html><html lang="tr"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">'
            f"<title>SSS önizleme · {html.escape(ad)}</title></head>"
            '<body style="margin:0; background:#0d0d0d; font-family:-apple-system, BlinkMacSystemFont, \'Segoe UI\', Roboto, sans-serif;">'
            '<div style="max-width:980px; margin:0 auto;">' + sss + "</div></body></html>")
    print(f"yazıldı: {a.out}\n  URL: {url}\n  Meta title ({len(title)}): {title}\n  Meta description ({len(desc)}): {desc}\n"
          f"  Uzun açıklama: {uzun.count('<h2>')} H2, {uzun.count('<h3>')} H3, {uzun.count('<a href')} link · FAQ: {len(d['sss'])} soru")


if __name__ == "__main__":
    main()
