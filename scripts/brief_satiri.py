#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Brief Excel'ine bir oyun satırı ekler. Şablonun biçimini ve sütun düzenini korur.

Kullanım:
    python3 brief_satiri.py --xlsx "GeForce NOW oyun detay sayfası içerik briefleri.xlsx" \
        --json satir.json

satir.json biçimi (dokuz alanın hepsi zorunlu, çok satırlı metinler "\n" ile):
{
  "oyun": "Battlefield 6",
  "main_kw": "battlefield 6",
  "hacim": 49500,
  "ikincil": "kelime (hacim)\nkelime (hacim)",
  "basliklar": "H2: ...\nH2: ...",
  "kurgu": "TON: ...\n\nAÇILIŞ: ...",
  "kurgu_kalin": "Koşullu uyarı cümlesi (isteğe bağlı, kalın basılır)",
  "link": "1. anchor : https://...\n   yerleşeceği bölüm",
  "sss": "1. Soru?\n   Yanıtta: ...",
  "yanit": "• Her yanıt ort. 20-60 kelimedir."
}
Dosya yoksa şablon başlık satırıyla oluşturulur.
"""
import argparse, json, math, os, sys

INK, HEAD, FN = "FF10332F", "434343", "Calibri"
BASLIK = ["Oyun Adı", "Main KW", "Main KW Hacim", "İkincil Kelimeler", "Alt Başlıklar",
          "İçerik Kurgusu", "Link Verilecek Sayfalar", "SSS'ler", "Yanıt Biçimi"]
GENISLIK = [22, 20, 12, 38, 48, 134, 72, 66, 58]
ANAHTAR = ["oyun", "main_kw", "hacim", "ikincil", "basliklar", "kurgu", "link", "sss", "yanit"]


def satir_sayisi(metin, genislik):
    kap = max(1, int(genislik * 1.15))
    return sum(max(1, math.ceil(len(p) / kap)) for p in str(metin).split("\n"))


def main():
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont

    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", required=True)
    ap.add_argument("--json", required=True)
    a = ap.parse_args()
    d = json.load(open(a.json, encoding="utf-8"))
    eksik = [k for k in ANAHTAR if k not in d]
    if eksik:
        sys.exit(f"satir.json içinde eksik alan: {', '.join(eksik)}")

    if os.path.exists(a.xlsx):
        wb = load_workbook(a.xlsx, rich_text=True)
        ws = wb.active
    else:
        wb = Workbook(); ws = wb.active; ws.title = "Oyun Detay Briefleri"
        for c, b in enumerate(BASLIK, start=1):
            h = ws.cell(row=1, column=c, value=b)
            h.font = Font(name=FN, size=11, bold=True, color="FFFFFF")
            h.fill = PatternFill("solid", fgColor=HEAD)
            h.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[1].height = 36
        for i, w in enumerate(GENISLIK, start=1):
            ws.column_dimensions[chr(ord("A") + i - 1)].width = w
        ws.freeze_panes = "A2"

    r = ws.max_row + 1
    if ws.cell(row=ws.max_row, column=1).value in (None, ""):
        r = ws.max_row
    for c, k in enumerate(ANAHTAR, start=1):
        deger = d[k]
        if k == "kurgu" and d.get("kurgu_kalin"):
            deger = CellRichText([
                TextBlock(InlineFont(rFont=FN, sz=10, b=False, color=INK), str(d["kurgu"])),
                TextBlock(InlineFont(rFont=FN, sz=10, b=True, color=INK), str(d["kurgu_kalin"]))])
        cell = ws.cell(row=r, column=c, value=deger)
        cell.font = Font(name=FN, size=10, color=INK)
        cell.alignment = Alignment(horizontal="center" if k == "hacim" else "left",
                                   vertical="center", wrap_text=True)
        if k == "hacim":
            cell.number_format = "#,##0"

    metinler = [str(d[k]) + str(d.get("kurgu_kalin", "") if k == "kurgu" else "") for k in ANAHTAR]
    en = max(satir_sayisi(m, GENISLIK[i]) for i, m in enumerate(metinler))
    ws.row_dimensions[r].height = min(409, max(30, round(en * 13.2)))
    wb.save(a.xlsx)
    print(f"{a.xlsx} · {r}. satıra '{d['oyun']}' eklendi (yükseklik {ws.row_dimensions[r].height})")
    if en * 13.2 > 409:
        print("UYARI: en uzun hücre 409 punto satır sınırını aşıyor, metnin tamamı ekranda görünmeyebilir.")


if __name__ == "__main__":
    main()
